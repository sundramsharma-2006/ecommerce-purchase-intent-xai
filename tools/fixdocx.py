"""Post-process pandoc output: fixed-width tables that fill the page, small table font, A4 margins.
python tools/fixdocx.py in.docx"""
import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

path = sys.argv[1]
d = Document(path)
for s in d.sections:
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(2.0)
usable = Cm(17.0)
for t in d.tables:
    t.autofit = False
    tblPr = t._tbl.tblPr
    for el in tblPr.findall(qn("w:tblLayout")): tblPr.remove(el)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    for el in tblPr.findall(qn("w:tblW")): tblPr.remove(el)
    w = OxmlElement("w:tblW"); w.set(qn("w:w"), str(int(usable.twips))); w.set(qn("w:type"), "dxa"); tblPr.append(w)
    ncol = len(t.rows[0].cells)
    if ncol == 0:
        continue
    first = Cm(4.2) if ncol >= 5 else usable / ncol
    widths = [first] + [int((usable - first) / (ncol - 1))] * (ncol - 1) if ncol > 1 else [usable]
    grid = t._tbl.tblGrid
    for el in list(grid):
        grid.remove(el)
    for wd in widths:
        gc = OxmlElement("w:gridCol"); gc.set(qn("w:w"), str(int(wd) // 635)); grid.append(gc)
    for row in t.rows:
        for c, wd in zip(row.cells, widths):
            c.width = wd
            for p in c.paragraphs:
                for r in p.runs: r.font.size = Pt(8.5)
    borders = OxmlElement("w:tblBorders")
    for b in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{b}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "BFC5D0"); borders.append(e)
    for el in tblPr.findall(qn("w:tblBorders")): tblPr.remove(el)
    tblPr.append(borders)
    for c in t.rows[0].cells:
        for p in c.paragraphs:
            for r in p.runs: r.font.bold = True
d.save(path)
print("fixed", len(d.tables), "tables")
